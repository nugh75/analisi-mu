"""
Modelli del database per l'applicazione di etichettatura tematica.
"""

from flask_sqlalchemy import SQLAlchemy
from flask_login import UserMixin
from werkzeug.security import generate_password_hash, check_password_hash
from datetime import datetime
import json

# Questa istanza db sarà importata da app.py
db = SQLAlchemy()

class User(UserMixin, db.Model):
    """Modello per gli utenti/etichettatori"""
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(20), default='annotatore')  # 'amministratore' o 'annotatore'
    is_active = db.Column(db.Boolean, default=True)
    last_login = db.Column(db.DateTime)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    annotations = db.relationship('CellAnnotation', foreign_keys='CellAnnotation.user_id', backref='user', lazy=True)
    
    def set_password(self, password):
        """Imposta la password hashata"""
        self.password_hash = generate_password_hash(password)
    
    def check_password(self, password):
        """Verifica la password"""
        return check_password_hash(self.password_hash, password)
    
    @property
    def is_admin(self):
        """Compatibilità con il codice esistente"""
        return self.role == 'amministratore'
    
    def can_manage_users(self):
        """Verifica se l'utente può gestire altri utenti"""
        return self.role == 'amministratore'
    
    def can_access_backup(self):
        """Verifica se l'utente può accedere al backup"""
        return self.role == 'amministratore'
    
    def update_last_login(self):
        """Aggiorna il timestamp dell'ultimo login"""
        self.last_login = datetime.utcnow()
        db.session.commit()
    
    def __repr__(self):
        return f'<User {self.username} ({self.role})>'

class Category(db.Model):
    """Modello per le categorie di etichette"""
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False, unique=True)
    description = db.Column(db.Text)
    color = db.Column(db.String(7), default='#6c757d')  # Colore HEX
    is_active = db.Column(db.Boolean, default=True)  # Soft delete
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    labels = db.relationship('Label', foreign_keys='Label.category_id', back_populates='category_obj')
    
    def __repr__(self):
        return f'<Category {self.name}>'
    
    @classmethod
    def assign_next_color(cls):
        """
        Assegna automaticamente il prossimo colore disponibile dalla palette
        
        Returns:
            str: Codice esadecimale del colore assegnato
        """
        from utils.color_palette import ColorPalette
        
        # Ottieni tutti i colori già utilizzati
        used_colors = [cat.color for cat in cls.query.all() if cat.color]
        
        return ColorPalette.get_next_color(used_colors)
    
    def get_effective_color(self):
        """
        Restituisce il colore effettivo della categoria
        
        Returns:
            str: Codice esadecimale del colore
        """
        return self.color or '#6c757d'
    
    def get_text_color(self):
        """
        Restituisce il colore del testo ottimale per questa categoria
        
        Returns:
            str: '#ffffff' per testo bianco, '#000000' per testo nero
        """
        from utils.color_palette import ColorPalette
        return ColorPalette.get_contrasting_text_color(self.get_effective_color())
    
    def update_color(self, new_color):
        """
        Aggiorna il colore della categoria e propaga a tutte le etichette
        che non hanno un colore personalizzato
        
        Args:
            new_color (str): Nuovo colore in formato esadecimale
        """
        from utils.color_palette import ColorPalette
        
        if not ColorPalette.validate_color(new_color):
            raise ValueError(f"Colore non valido: {new_color}")
        
        # Aggiorna il colore della categoria
        self.color = new_color
        
        # Propaga il nuovo colore a tutte le etichette che non hanno un colore personalizzato
        self.propagate_color_to_labels(force=False)
    
    def propagate_color_to_labels(self, force=False):
        """
        Propaga il colore della categoria a tutte le sue etichette
        
        Args:
            force (bool): Se True, sovrascrive anche i colori personalizzati
        """
        for label in self.labels:
            if force or not label.has_custom_color():
                label.color = self.color
    
    def get_hsl(self):
        """
        Restituisce i valori HSL del colore della categoria
        
        Returns:
            tuple: (hue, saturation, lightness) dove:
                   - hue è tra 0-360
                   - saturation è tra 0-1
                   - lightness è tra 0-1
        """
        from utils.color_palette import ColorPalette
        return ColorPalette.hex_to_hsl(self.get_effective_color())
    
    def set_hsl(self, hue, saturation, lightness):
        """
        Imposta il colore della categoria usando valori HSL
        
        Args:
            hue (float): Tonalità (0-360)
            saturation (float): Saturazione (0-1)
            lightness (float): Luminosità (0-1)
        """
        from utils.color_palette import ColorPalette
        hex_color = ColorPalette.hsl_to_hex(hue, saturation, lightness)
        self.update_color(hex_color)

class Label(db.Model):
    """Modello per le etichette globali"""
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False, unique=True)
    description = db.Column(db.Text)
    category_id = db.Column(db.Integer, db.ForeignKey('category.id'), nullable=True)
    category = db.Column(db.String(50))  # Manteniamo per compatibilità, sarà deprecato
    color = db.Column(db.String(7), default='#007bff')  # Colore HEX
    is_active = db.Column(db.Boolean, default=True)  # Soft delete
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    annotations = db.relationship('CellAnnotation', backref='label', lazy=True, cascade='all, delete-orphan')
    category_obj = db.relationship('Category', foreign_keys=[category_id], back_populates='labels')
    
    def __repr__(self):
        return f'<Label {self.name}>'
    
    def get_effective_color(self):
        """
        Restituisce il colore effettivo dell'etichetta
        Se l'etichetta ha una categoria, usa il colore della categoria
        altrimenti usa il colore proprio dell'etichetta
        
        Returns:
            str: Codice esadecimale del colore
        """
        if self.category_obj and not self.has_custom_color():
            return self.category_obj.get_effective_color()
        return self.color or '#007bff'
    
    def has_custom_color(self):
        """
        Verifica se l'etichetta ha un colore personalizzato diverso da quello della categoria
        
        Returns:
            bool: True se ha un colore personalizzato
        """
        if not self.category_obj:
            return True  # Se non ha categoria, il colore è sempre personalizzato
        
        return self.color != self.category_obj.color
    
    def inherit_category_color(self):
        """
        Fa ereditare il colore dalla categoria (rimuove personalizzazione)
        """
        if self.category_obj:
            self.color = self.category_obj.color
    
    def get_text_color(self):
        """
        Restituisce il colore del testo ottimale per questa etichetta
        
        Returns:
            str: '#ffffff' per testo bianco, '#000000' per testo nero
        """
        from utils.color_palette import ColorPalette
        return ColorPalette.get_contrasting_text_color(self.get_effective_color())
    
    def set_color(self, new_color):
        """
        Imposta un nuovo colore per l'etichetta
        
        Args:
            new_color (str): Nuovo colore in formato esadecimale
        """
        from utils.color_palette import ColorPalette
        
        if not ColorPalette.validate_color(new_color):
            raise ValueError(f"Colore non valido: {new_color}")
        
        self.color = new_color

class ExcelFile(db.Model):
    """Modello per i file Excel caricati"""
    id = db.Column(db.Integer, primary_key=True)
    filename = db.Column(db.String(255), nullable=False)
    original_filename = db.Column(db.String(255), nullable=False)
    file_path = db.Column(db.String(500), nullable=False)
    uploaded_by = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    uploaded_at = db.Column(db.DateTime, default=datetime.utcnow)
    
    # Relazioni
    uploader = db.relationship('User', backref='uploaded_files')
    text_cells = db.relationship('TextCell', backref='excel_file', lazy=True, cascade='all, delete-orphan')
    
    def __repr__(self):
        return f'<ExcelFile {self.original_filename}>'

class TextCell(db.Model):
    """Modello per le celle testuali estratte dai file Excel"""
    id = db.Column(db.Integer, primary_key=True)
    excel_file_id = db.Column(db.Integer, db.ForeignKey('excel_file.id'), nullable=False)
    sheet_name = db.Column(db.String(100), nullable=False)
    row_index = db.Column(db.Integer, nullable=False)
    column_index = db.Column(db.Integer, nullable=False)
    column_name = db.Column(db.String(100))
    text_content = db.Column(db.Text, nullable=False)
    
    # Classificazione manuale del tipo di domanda
    question_type = db.Column(db.String(20))  # 'aperta', 'chiusa_binaria', 'chiusa_multipla', 'likert', 'anagrafica', 'numerica'
    
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    
    # Relazioni
    # annotations verranno creati automaticamente dal backref in CellAnnotation
    
    @property
    def cell_reference(self):
        """Restituisce il riferimento della cella (es. A1, B2)"""
        from openpyxl.utils import get_column_letter
        return f"{get_column_letter(self.column_index + 1)}{self.row_index + 1}"
    
    def __repr__(self):
        return f'<TextCell {self.sheet_name}:{self.cell_reference}>'

class AIConfiguration(db.Model):
    """Configurazione per i provider AI"""
    id = db.Column(db.Integer, primary_key=True)
    provider = db.Column(db.String(20), nullable=False)  # 'ollama' o 'openrouter'
    name = db.Column(db.String(100), nullable=False)  # Nome descrittivo
    
    # Configurazione Ollama
    ollama_url = db.Column(db.String(255))
    ollama_model = db.Column(db.String(100))
    
    # Configurazione OpenRouter
    openrouter_api_key = db.Column(db.String(255))
    openrouter_model = db.Column(db.String(100))
    
    # Configurazione generale
    is_active = db.Column(db.Boolean, default=True)
    max_tokens = db.Column(db.Integer, default=1000)
    temperature = db.Column(db.Float, default=0.7)
    system_prompt = db.Column(db.Text)
    
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    def to_dict(self):
        return {
            'id': self.id,
            'provider': self.provider,
            'name': self.name,
            'ollama_url': self.ollama_url,
            'ollama_model': self.ollama_model,
            'openrouter_model': self.openrouter_model,
            'is_active': self.is_active,
            'max_tokens': self.max_tokens,
            'temperature': self.temperature
        }

class OpenRouterModel(db.Model):
    """Modelli disponibili su OpenRouter"""
    id = db.Column(db.Integer, primary_key=True)
    model_id = db.Column(db.String(100), nullable=False, unique=True)
    name = db.Column(db.String(200), nullable=False)
    description = db.Column(db.Text)
    context_length = db.Column(db.Integer)
    pricing_prompt = db.Column(db.Float)  # Prezzo per token di input
    pricing_completion = db.Column(db.Float)  # Prezzo per token di output
    is_free = db.Column(db.Boolean, default=False)
    is_available = db.Column(db.Boolean, default=True)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

class OllamaModel(db.Model):
    """Modelli disponibili su Ollama"""
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    tag = db.Column(db.String(50), nullable=False)
    size = db.Column(db.BigInteger)  # Dimensione in bytes
    digest = db.Column(db.String(100))
    modified_at = db.Column(db.DateTime)
    is_pulled = db.Column(db.Boolean, default=False)
    pull_progress = db.Column(db.Float, default=0.0)  # Progresso del download 0-100
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    __table_args__ = (db.UniqueConstraint('name', 'tag', name='unique_model_tag'),)

class CellAnnotation(db.Model):
    """Modello per le annotazioni delle celle"""
    id = db.Column(db.Integer, primary_key=True)
    text_cell_id = db.Column(db.Integer, db.ForeignKey('text_cell.id'), nullable=False)
    label_id = db.Column(db.Integer, db.ForeignKey('label.id'), nullable=False)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Campi specifici per AI
    is_ai_generated = db.Column(db.Boolean, default=False)
    ai_confidence = db.Column(db.Float)  # Confidence score 0-1
    ai_model = db.Column(db.String(100))  # Modello AI utilizzato
    ai_provider = db.Column(db.String(20))  # Provider AI (ollama/openrouter)
    status = db.Column(db.String(20), default='active')  # active, pending_review, rejected
    reviewed_by = db.Column(db.Integer, db.ForeignKey('user.id'))  # Chi ha approvato/rifiutato
    reviewed_at = db.Column(db.DateTime)
    
    # Relazioni
    text_cell = db.relationship('TextCell', backref='annotations')
    reviewer = db.relationship('User', foreign_keys=[reviewed_by])
    
    def to_dict(self):
        return {
            'id': self.id,
            'text_cell_id': self.text_cell_id,
            'label_id': self.label_id,
            'user_id': self.user_id,
            'is_ai_generated': self.is_ai_generated,
            'ai_confidence': self.ai_confidence,
            'ai_model': self.ai_model,
            'ai_provider': self.ai_provider,
            'status': self.status,
            'created_at': self.created_at.isoformat() if self.created_at else None,
            'reviewed_at': self.reviewed_at.isoformat() if self.reviewed_at else None
        }
    
    def __repr__(self):
        return f'<CellAnnotation {self.id}: Cell {self.text_cell_id} -> Label {self.label_id}>'

class AnnotationAction(db.Model):
    """Modello per tracciare le azioni sulle annotazioni"""
    id = db.Column(db.Integer, primary_key=True)
    text_cell_id = db.Column(db.Integer, db.ForeignKey('text_cell.id'), nullable=False)
    label_id = db.Column(db.Integer, db.ForeignKey('label.id'), nullable=False)
    action_type = db.Column(db.String(20), nullable=False)  # 'added', 'removed', 'approved', 'rejected'
    performed_by = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)  # Chi ha eseguito l'azione
    target_user_id = db.Column(db.Integer, db.ForeignKey('user.id'))  # Su quale annotazione utente (può essere diverso da performed_by)
    annotation_id = db.Column(db.Integer, db.ForeignKey('cell_annotation.id'))  # L'annotazione interessata (se esiste ancora)
    notes = db.Column(db.Text)  # Note aggiuntive
    timestamp = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    
    # Campi per tracciare annotazioni AI
    was_ai_generated = db.Column(db.Boolean, default=False)
    ai_confidence = db.Column(db.Float)
    ai_model = db.Column(db.String(100))
    ai_provider = db.Column(db.String(20))
    
    # Relazioni
    text_cell = db.relationship('TextCell', foreign_keys=[text_cell_id])
    label = db.relationship('Label', foreign_keys=[label_id])
    performer = db.relationship('User', foreign_keys=[performed_by])
    target_user = db.relationship('User', foreign_keys=[target_user_id])
    annotation = db.relationship('CellAnnotation', foreign_keys=[annotation_id])
    
    def __repr__(self):
        return f'<AnnotationAction {self.id}: {self.action_type} on Cell {self.text_cell_id} by User {self.performed_by}>'

class AIPromptTemplate(db.Model):
    """Template per prompt AI dinamici"""
    __tablename__ = 'ai_prompt_template'
    
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False, unique=True)
    description = db.Column(db.Text)
    category = db.Column(db.String(100))  # Categoria del template
    base_prompt = db.Column(db.Text, nullable=False)  # Parte generica del prompt
    is_active = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    def build_prompt_with_categories(self, selected_categories, labels, texts):
        """Costruisce il prompt finale con le categorie selezionate"""
        from collections import defaultdict
        
        # Filtra le etichette per le categorie selezionate
        filtered_labels = [
            label for label in labels 
            if label.category_obj and label.category_obj.name in selected_categories
        ]
        
        if not filtered_labels:
            # Se nessuna categoria, usa tutte le etichette
            filtered_labels = labels
        
        # Organizza per categoria
        categories_dict = defaultdict(list)
        for label in filtered_labels:
            if label.is_active:
                category_name = label.category_obj.name if label.category_obj else 'Generale'
                categories_dict[category_name].append(label)
        
        # Costruisce la sezione etichette
        categories_text = "ETICHETTE DISPONIBILI (categorie selezionate):\n"
        for category, cat_labels in categories_dict.items():
            categories_text += f"\n=== {category} ===\n"
            for label in cat_labels:
                desc = f" - {label.description}" if label.description else ""
                categories_text += f"• {label.name}{desc}\n"
        
        # Costruisce la sezione testi
        texts_section = "\nTESTI DA ANALIZZARE:\n"
        for i, text in enumerate(texts):
            texts_section += f"{i}. {text}\n"
        
        # Combina tutto
        full_prompt = f"{self.base_prompt}\n\n{categories_text}\n{texts_section}"
        
        return full_prompt
    
    def __repr__(self):
        return f'<AIPromptTemplate {self.name}>'

class PromptTemplate(db.Model):
    """Template di prompt per analisi AI categorizzate"""
    __tablename__ = 'prompt_templates'
    
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    category = db.Column(db.String(100), nullable=False)
    description = db.Column(db.Text)
    template_text = db.Column(db.Text, nullable=False)
    is_active = db.Column(db.Boolean, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    def __repr__(self):
        return f'<PromptTemplate {self.name} ({self.category})>'


class TextDocument(db.Model):
    """Modello per documenti di testo (focus group, interviste, etc.)"""
    __tablename__ = 'text_documents'
    
    id = db.Column(db.Integer, primary_key=True)
    filename = db.Column(db.String(255), nullable=False)  # Nome file sul server
    original_name = db.Column(db.String(255), nullable=False)  # Nome originale del file
    content = db.Column(db.Text, nullable=False)  # Contenuto del documento
    document_type = db.Column(db.String(50), default='other')  # 'focus_group', 'interview', 'other'
    file_format = db.Column(db.String(10), nullable=False)  # 'txt', 'md', 'docx'
    
    # Metadati
    word_count = db.Column(db.Integer, default=0)
    character_count = db.Column(db.Integer, default=0)
    
    # Relazioni utente
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    
    # Timestamp
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    uploader = db.relationship('User', backref='text_documents')
    annotations = db.relationship('TextAnnotation', backref='document', lazy=True, cascade='all, delete-orphan')
    
    def __repr__(self):
        return f'<TextDocument {self.original_name}>'
    
    def update_stats(self):
        """Aggiorna le statistiche del documento"""
        if self.content:
            self.word_count = len(self.content.split())
            self.character_count = len(self.content)


class TextAnnotation(db.Model):
    """Modello per le annotazioni su documenti di testo"""
    __tablename__ = 'text_annotations'
    
    id = db.Column(db.Integer, primary_key=True)
    document_id = db.Column(db.Integer, db.ForeignKey('text_documents.id'), nullable=False)
    
    # Selezione di testo
    text_selection = db.Column(db.Text, nullable=False)  # Testo selezionato
    start_position = db.Column(db.Integer, nullable=False)  # Posizione inizio
    end_position = db.Column(db.Integer, nullable=False)  # Posizione fine
    
    # Contesto (righe prima e dopo per riferimento)
    context_before = db.Column(db.Text)
    context_after = db.Column(db.Text)
    
    # Etichetta assegnata
    label_id = db.Column(db.Integer, db.ForeignKey('label.id'), nullable=False)
    
    # Utente che ha creato l'annotazione
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    
    # Metadati AI (se generata automaticamente)
    is_ai_generated = db.Column(db.Boolean, default=False)
    ai_confidence = db.Column(db.Float)
    ai_model = db.Column(db.String(100))
    ai_provider = db.Column(db.String(50))
    
    # Status e review
    status = db.Column(db.String(20), default='active')  # 'active', 'disputed', 'archived'
    reviewed_by = db.Column(db.Integer, db.ForeignKey('user.id'))
    reviewed_at = db.Column(db.DateTime)
    
    # Timestamp
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    label = db.relationship('Label', backref='text_annotations')
    user = db.relationship('User', foreign_keys=[user_id], backref='text_annotations')
    reviewer = db.relationship('User', foreign_keys=[reviewed_by])
    
    def __repr__(self):
        return f'<TextAnnotation {self.text_selection[:50]}...>'
    
    @property
    def text_preview(self):
        """Anteprima del testo selezionato (massimo 100 caratteri)"""
        if len(self.text_selection) <= 100:
            return self.text_selection
        return f"{self.text_selection[:97]}..."


class ForumCategory(db.Model):
    """Modello per le categorie del forum"""
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    description = db.Column(db.Text)
    icon = db.Column(db.String(50), default='bi-chat-dots')
    color = db.Column(db.String(20), default='primary')
    is_system = db.Column(db.Boolean, default=False)  # True per categorie prestabilite (domande)
    excel_file_id = db.Column(db.Integer, db.ForeignKey('excel_file.id'), nullable=False)
    question_name = db.Column(db.String(255))  # Solo per categorie di domande
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    created_by = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    
    # Relazioni
    excel_file = db.relationship('ExcelFile', backref='forum_categories')
    creator = db.relationship('User', backref='created_categories')
    posts = db.relationship('ForumPost', backref='category', lazy=True, cascade='all, delete-orphan')
    
    def __repr__(self):
        return f'<ForumCategory {self.name}>'
    
    @property
    def post_count(self):
        """Numero di post nella categoria"""
        return len(self.posts)
    
    @property
    def last_post(self):
        """Ultimo post nella categoria"""
        return ForumPost.query.filter_by(category_id=self.id).order_by(ForumPost.created_at.desc()).first()


class ForumPost(db.Model):
    """Modello per i post del forum"""
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    content = db.Column(db.Text, nullable=False)
    is_pinned = db.Column(db.Boolean, default=False)
    is_locked = db.Column(db.Boolean, default=False)
    views = db.Column(db.Integer, default=0)
    
    # Relazioni
    category_id = db.Column(db.Integer, db.ForeignKey('forum_category.id'), nullable=False)
    author_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    
    # Timestamp
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    author = db.relationship('User', backref='forum_posts')
    comments = db.relationship('ForumComment', backref='post', lazy=True, cascade='all, delete-orphan')
    
    def __repr__(self):
        return f'<ForumPost {self.title}>'
    
    @property
    def comment_count(self):
        """Numero di commenti nel post"""
        return len(self.comments)
    
    @property
    def last_comment(self):
        """Ultimo commento nel post"""
        return ForumComment.query.filter_by(post_id=self.id).order_by(ForumComment.created_at.desc()).first()
    
    def increment_views(self):
        """Incrementa il contatore delle visualizzazioni"""
        self.views += 1
        db.session.commit()


class ForumComment(db.Model):
    """Modello per i commenti dei post del forum"""
    id = db.Column(db.Integer, primary_key=True)
    content = db.Column(db.Text, nullable=False)
    
    # Relazioni
    post_id = db.Column(db.Integer, db.ForeignKey('forum_post.id'), nullable=False)
    author_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    parent_id = db.Column(db.Integer, db.ForeignKey('forum_comment.id'))  # Per reply annidate
    
    # Timestamp
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    author = db.relationship('User', backref='forum_comments')
    parent = db.relationship('ForumComment', remote_side=[id], backref='replies')
    
    def __repr__(self):
        return f'<ForumComment by {self.author.username}>'
    
    @property
    def content_preview(self):
        """Anteprima del contenuto (massimo 100 caratteri)"""
        if len(self.content) <= 100:
            return self.content
        return f"{self.content[:97]}..."


# Modelli per il sistema di decisioni sull'etichettatura
class DecisionSession(db.Model):
    """Sessione di decisione per il raggruppamento delle etichette"""
    __tablename__ = 'decision_session'
    
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    description = db.Column(db.Text)
    status = db.Column(db.String(20), default='draft')  # draft, active, completed, archived
    
    # Configurazioni della sessione
    voting_threshold = db.Column(db.Float, default=0.6)  # Soglia per approvazione (60%)
    allow_public_comments = db.Column(db.Boolean, default=True)
    
    # Metadati
    created_by = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    creator = db.relationship('User', backref='created_decision_sessions')
    proposals = db.relationship('LabelGroupingProposal', backref='session', lazy=True, cascade='all, delete-orphan')
    comments = db.relationship('LabelDecisionComment', backref='session', lazy=True, cascade='all, delete-orphan')
    
    def __repr__(self):
        return f'<DecisionSession {self.name}>'
    
    @property
    def proposal_count(self):
        """Numero totale di proposte"""
        return len(self.proposals)
    
    @property
    def approved_proposals_count(self):
        """Numero di proposte approvate"""
        return len([p for p in self.proposals if p.status == 'approved'])
    
    @property
    def pending_proposals_count(self):
        """Numero di proposte in attesa"""
        return len([p for p in self.proposals if p.status == 'pending'])
    
    @property
    def completion_percentage(self):
        """Percentuale di completamento della sessione"""
        if self.proposal_count == 0:
            return 0
        return (self.approved_proposals_count / self.proposal_count) * 100
    
    def can_edit(self, user):
        """Verifica se l'utente può modificare la sessione"""
        return user.is_admin or user.id == self.created_by
    
    def get_participants(self):
        """Ottiene la lista degli utenti che hanno partecipato"""
        participant_ids = set()
        for proposal in self.proposals:
            participant_ids.add(proposal.created_by)
            for vote in proposal.votes:
                participant_ids.add(vote.user_id)
        return User.query.filter(User.id.in_(participant_ids)).all()


class LabelGroupingProposal(db.Model):
    """Proposta di raggruppamento delle etichette"""
    __tablename__ = 'label_grouping_proposal'
    
    id = db.Column(db.Integer, primary_key=True)
    session_id = db.Column(db.Integer, db.ForeignKey('decision_session.id'), nullable=False)
    
    # Contenuto della proposta
    category = db.Column(db.String(200), nullable=False)  # es: "Impatti sugli studenti"
    original_labels = db.Column(db.Text, nullable=False)  # JSON array delle etichette originali
    proposed_label = db.Column(db.String(200), nullable=False)  # es: "Accessibilità e inclusione"
    proposed_code = db.Column(db.String(10), nullable=False)  # es: "S1"
    rationale = db.Column(db.Text)  # Motivazione del raggruppamento
    
    # Status e metadati
    status = db.Column(db.String(20), default='pending')  # pending, approved, rejected, under_review
    created_by = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    creator = db.relationship('User', backref='created_proposals')
    votes = db.relationship('LabelDecisionVote', backref='proposal', lazy=True, cascade='all, delete-orphan')
    comments = db.relationship('LabelDecisionComment', backref='proposal', lazy=True, cascade='all, delete-orphan')
    
    def __repr__(self):
        return f'<LabelGroupingProposal {self.proposed_code}: {self.proposed_label}>'
    
    @property
    def original_labels_list(self):
        """Ottiene la lista delle etichette originali"""
        try:
            return json.loads(self.original_labels)
        except:
            return self.original_labels.split(';') if self.original_labels else []
    
    @original_labels_list.setter
    def original_labels_list(self, labels):
        """Imposta la lista delle etichette originali"""
        if isinstance(labels, list):
            self.original_labels = json.dumps(labels)
        else:
            self.original_labels = labels
    
    @property
    def vote_counts(self):
        """Conta i voti per tipo"""
        counts = {'approve': 0, 'reject': 0, 'abstain': 0}
        for vote in self.votes:
            counts[vote.vote] = counts.get(vote.vote, 0) + 1
        return counts
    
    @property
    def total_votes(self):
        """Numero totale di voti"""
        return len(self.votes)
    
    @property
    def approval_percentage(self):
        """Percentuale di approvazione"""
        if self.total_votes == 0:
            return 0
        return (self.vote_counts['approve'] / self.total_votes) * 100
    
    def get_user_vote(self, user_id):
        """Ottiene il voto di un utente specifico"""
        for vote in self.votes:
            if vote.user_id == user_id:
                return vote
        return None
    
    def update_status_by_votes(self, threshold=0.6):
        """Aggiorna lo status basato sui voti"""
        if self.total_votes < 2:  # Minimo 2 voti per decidere
            return
        
        if self.approval_percentage >= (threshold * 100):
            self.status = 'approved'
        elif self.vote_counts['reject'] > self.vote_counts['approve']:
            self.status = 'rejected'


class LabelDecisionVote(db.Model):
    """Voto su una proposta di raggruppamento"""
    __tablename__ = 'label_decision_vote'
    
    id = db.Column(db.Integer, primary_key=True)
    proposal_id = db.Column(db.Integer, db.ForeignKey('label_grouping_proposal.id'), nullable=False)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    
    # Voto e commento
    vote = db.Column(db.String(10), nullable=False)  # approve, reject, abstain
    comment = db.Column(db.Text)  # Commento opzionale per motivare il voto
    
    # Timestamp
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    user = db.relationship('User', backref='decision_votes')
    
    # Constraint per un voto per utente per proposta
    __table_args__ = (db.UniqueConstraint('proposal_id', 'user_id', name='unique_vote_per_user_proposal'),)
    
    def __repr__(self):
        return f'<LabelDecisionVote {self.user.username}: {self.vote}>'


class LabelDecisionComment(db.Model):
    """Commento su una sessione di decisione o proposta specifica"""
    __tablename__ = 'label_decision_comment'
    
    id = db.Column(db.Integer, primary_key=True)
    session_id = db.Column(db.Integer, db.ForeignKey('decision_session.id'), nullable=False)
    proposal_id = db.Column(db.Integer, db.ForeignKey('label_grouping_proposal.id'), nullable=True)  # Opzionale
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    
    # Contenuto
    content = db.Column(db.Text, nullable=False)
    parent_id = db.Column(db.Integer, db.ForeignKey('label_decision_comment.id'))  # Per reply annidate
    
    # Timestamp
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    user = db.relationship('User', backref='decision_comments')
    parent = db.relationship('LabelDecisionComment', remote_side=[id], backref='replies')
    
    def __repr__(self):
        return f'<LabelDecisionComment by {self.user.username}>'
    
    @property
    def content_preview(self):
        """Anteprima del contenuto (massimo 100 caratteri)"""
        if len(self.content) <= 100:
            return self.content
        return f"{self.content[:97]}..."
    
    @property
    def is_general_comment(self):
        """True se è un commento generale sulla sessione"""
        return self.proposal_id is None


class DiaryEntry(db.Model):
    """Modello per le voci del Diario di Bordo"""
    __tablename__ = 'diary_entries'
    
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    content = db.Column(db.Text, nullable=False)
    activity_type = db.Column(db.String(50), default='general')  # general, meeting, milestone, issue, decision, reflection, consideration
    priority = db.Column(db.String(20), default='medium')  # low, medium, high, urgent
    status = db.Column(db.String(20), default='active')  # active, completed, archived
    
    # Data/ora personalizzabile per la voce
    entry_date = db.Column(db.DateTime, default=datetime.utcnow)  # Data dell'attività/evento
    
    # Metadati per export
    word_count = db.Column(db.Integer, default=0)
    mentioned_users = db.Column(db.Text)  # JSON array degli utenti menzionati
    referenced_files = db.Column(db.Text)  # JSON array dei file referenziati
    tags = db.Column(db.Text)  # JSON array dei tag
    
    # Relazioni
    author_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    project_id = db.Column(db.Integer, db.ForeignKey('excel_file.id'), nullable=True)  # Progetto correlato
    
    # Timestamp
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    
    # Relazioni
    author = db.relationship('User', backref='diary_entries')
    project = db.relationship('ExcelFile', backref='diary_entries')
    attachments = db.relationship('DiaryAttachment', backref='entry', lazy=True, cascade='all, delete-orphan')
    
    def __repr__(self):
        return f'<DiaryEntry {self.title}>'
    
    @property
    def mentioned_users_list(self):
        """Ottiene la lista degli utenti menzionati"""
        try:
            return json.loads(self.mentioned_users) if self.mentioned_users else []
        except:
            return []
    
    @mentioned_users_list.setter
    def mentioned_users_list(self, users):
        """Imposta la lista degli utenti menzionati"""
        self.mentioned_users = json.dumps(users) if users else None
    
    @property
    def referenced_files_list(self):
        """Ottiene la lista dei file referenziati"""
        try:
            return json.loads(self.referenced_files) if self.referenced_files else []
        except:
            return []
    
    @referenced_files_list.setter
    def referenced_files_list(self, files):
        """Imposta la lista dei file referenziati"""
        self.referenced_files = json.dumps(files) if files else None
    
    @property
    def tags_list(self):
        """Ottiene la lista dei tag"""
        try:
            return json.loads(self.tags) if self.tags else []
        except:
            return []
    
    @tags_list.setter
    def tags_list(self, tags_list):
        """Imposta la lista dei tag"""
        self.tags = json.dumps(tags_list) if tags_list else None
    
    @property
    def content_preview(self):
        """Anteprima del contenuto (massimo 150 caratteri)"""
        if len(self.content) <= 150:
            return self.content
        return f"{self.content[:147]}..."
    
    def update_word_count(self):
        """Aggiorna il conteggio delle parole"""
        if self.content:
            # Rimuovi markdown e HTML per il conteggio
            import re
            clean_content = re.sub(r'[#*_`~\[\]()]', '', self.content)
            clean_content = re.sub(r'<[^>]+>', '', clean_content)
            self.word_count = len(clean_content.split())
    
    def parse_mentions_and_files(self):
        """Analizza il contenuto per estrarre menzioni e riferimenti"""
        import re
        
        # Estrai menzioni utenti (@username)
        user_mentions = re.findall(r'@(\w+)', self.content)
        
        # Estrai riferimenti file (#filename)
        file_references = re.findall(r'#(\S+)', self.content)
        
        # Aggiorna le liste
        self.mentioned_users_list = list(set(user_mentions))
        self.referenced_files_list = list(set(file_references))
    
    def get_activity_icon(self):
        """Restituisce l'icona Bootstrap per il tipo di attività"""
        icons = {
            'general': 'bi-journal-text',
            'meeting': 'bi-people',
            'milestone': 'bi-flag',
            'issue': 'bi-exclamation-triangle',
            'decision': 'bi-check-circle',
            'reflection': 'bi-lightbulb',
            'consideration': 'bi-chat-dots'
        }
        return icons.get(self.activity_type, 'bi-journal-text')
    
    def get_priority_class(self):
        """Restituisce la classe CSS per la priorità"""
        classes = {
            'low': 'text-success',
            'medium': 'text-warning',
            'high': 'text-danger',
            'urgent': 'text-danger fw-bold'
        }
        return classes.get(self.priority, 'text-secondary')


class DiaryAttachment(db.Model):
    """Modello per gli allegati del diario di bordo"""
    __tablename__ = 'diary_attachments'
    
    id = db.Column(db.Integer, primary_key=True)
    entry_id = db.Column(db.Integer, db.ForeignKey('diary_entries.id'), nullable=False)
    filename = db.Column(db.String(255), nullable=False)
    original_name = db.Column(db.String(255), nullable=False)
    file_path = db.Column(db.String(500), nullable=False)
    file_size = db.Column(db.Integer)  # Dimensione in bytes
    mime_type = db.Column(db.String(100))
    
    # Timestamp
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    
    def __repr__(self):
        return f'<DiaryAttachment {self.original_name}>'
    
    @property
    def file_size_human(self):
        """Dimensione del file in formato leggibile"""
        if not self.file_size:
            return "Sconosciuta"
        
        for unit in ['B', 'KB', 'MB', 'GB']:
            if self.file_size < 1024:
                return f"{self.file_size:.1f} {unit}"
            self.file_size /= 1024
        return f"{self.file_size:.1f} TB"
